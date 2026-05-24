#!/usr/bin/env python3
"""Ratio-Solvent Excel: X=phi1/(phi1+phi3), Y=phi2. Binodal + Spinodal. 6 systems in 1 sheet."""
import os, csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

DATA_DIR = os.path.join(os.path.dirname(__file__), '..', 'output', 'data')
OUT_PATH = os.path.join(os.path.dirname(__file__), 'Ratio_Solvent_All_Systems_v2.xlsx')

SYSTEMS = [
    ('PM6_Tol',  'PM6 / L8-Bo / Toluene',   '4472C4'),
    ('PM6_OXy',  'PM6 / L8-Bo / o-Xylene',  'ED7D31'),
    ('D18_Tol',  'D18 / L8-Bo / Toluene',   '70AD47'),
    ('D18_OXy',  'D18 / L8-Bo / o-Xylene',  'FFC000'),
    ('PM6_CF',   'PM6 / L8-Bo / Chloroform','5B9BD5'),
    ('D18_CF',   'D18 / L8-Bo / Chloroform', 'A5A5A5'),
]

COLS_PER_SYS = 5  # B_X, B_Y, S_X, S_Y, spacer

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Ratio-Solvent All'

# Row 1: Title
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COLS_PER_SYS*len(SYSTEMS))
ws['A1'] = 'Ratio-Solvent Phase Diagram: Binodal (concentrated) + Spinodal'
ws['A1'].font = Font(bold=True, size=14)

# Row 2: Note
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COLS_PER_SYS*len(SYSTEMS))
ws['A2'] = 'X = phi1/(phi1+phi3) = Acceptor/(Acceptor+Donor) | Y = phi2 = Solvent volume fraction'
ws['A2'].font = Font(color='666666', size=10)

# Row 3: System titles (merged across 4 data cols per system)
col = 1
for tag, title, color in SYSTEMS:
    ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col+3)
    cell = ws.cell(row=3, column=col, value=title)
    cell.font = Font(bold=True, size=12, color='FFFFFF')
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
    cell.alignment = Alignment(horizontal='center')
    col += COLS_PER_SYS

# Row 4: Sub-headers
col = 1
for tag, title, color in SYSTEMS:
    for sub, sub_color in [('Bino X','4472C4'), ('Bino Y','4472C4'), ('Spin X','ED7D31'), ('Spin Y','ED7D31')]:
        c = ws.cell(row=4, column=col, value=sub)
        c.font = Font(bold=True, color=sub_color, size=9)
        col += 1
    col += 1  # spacer

# Load all data
all_bino = {}
all_spin = {}
max_rows = 0

for tag, title, color in SYSTEMS:
    # --- BINODAL (concentrated phase only) ---
    bin_file = os.path.join(DATA_DIR, f'binodal_{tag}.csv')
    bino_xy = []
    if os.path.exists(bin_file):
        with open(bin_file, 'r') as f:
            reader = csv.reader(f)
            next(reader)
            rows = list(reader)
            for i in range(0, len(rows)-1, 2):  # odd rows = concentrated
                row = rows[i]
                if len(row) >= 4:
                    phi1, phi3, phi2 = float(row[1]), float(row[2]), float(row[3])
                    denom = phi1 + phi3
                    if denom > 0:
                        bino_xy.append((phi1/denom, phi2))
    bino_xy.sort(key=lambda p: p[0])

    # --- SPINODAL ---
    spin_file = os.path.join(DATA_DIR, f'spinodal_{tag}.csv')
    spin_xy = []
    if os.path.exists(spin_file):
        with open(spin_file, 'r') as f:
            reader = csv.reader(f)
            next(reader)
            for row in reader:
                if len(row) >= 4:
                    phi1, phi3, phi2 = float(row[1]), float(row[2]), float(row[3])
                    if 0 <= phi1 <= 1 and 0 <= phi2 <= 1 and 0 <= phi3 <= 1:
                        denom = phi1 + phi3
                        if denom > 0:
                            spin_xy.append((phi1/denom, phi2))
    spin_xy.sort(key=lambda p: p[0])

    all_bino[tag] = bino_xy
    all_spin[tag] = spin_xy
    max_rows = max(max_rows, len(bino_xy), len(spin_xy))

# Write data
for row_idx in range(max_rows):
    excel_row = row_idx + 5
    col = 1
    for tag, title, color in SYSTEMS:
        bino = all_bino[tag]
        spin = all_spin[tag]
        if row_idx < len(bino):
            ws.cell(row=excel_row, column=col,   value=round(bino[row_idx][0], 6))
            ws.cell(row=excel_row, column=col+1, value=round(bino[row_idx][1], 6))
        if row_idx < len(spin):
            ws.cell(row=excel_row, column=col+2, value=round(spin[row_idx][0], 6))
            ws.cell(row=excel_row, column=col+3, value=round(spin[row_idx][1], 6))
        col += COLS_PER_SYS

# Column widths
col = 1
for tag, title, color in SYSTEMS:
    for _ in range(4):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 12
        col += 1
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 2  # spacer
    col += 1

# Critical point summary
cp_row = max_rows + 7
ws.merge_cells(start_row=cp_row, start_column=1, end_row=cp_row, end_column=10)
ws.cell(row=cp_row, column=1, value='Critical Points').font = Font(bold=True, size=12)
cp_row += 1
for j, h in enumerate(['System','CP X','CP Y','phi1','phi2','phi3','Source']):
    ws.cell(row=cp_row, column=j+1, value=h).font = Font(bold=True)

for tag, title, color in SYSTEMS:
    crit_file = os.path.join(DATA_DIR, f'critical_{tag}.csv')
    if os.path.exists(crit_file):
        with open(crit_file, 'r') as f:
            reader = csv.reader(f); next(reader); row = next(reader)
            phi1, phi2, phi3 = float(row[0]), float(row[1]), float(row[2])
            cp_x = phi1/(phi1+phi3) if (phi1+phi3)>0 else 0
            cp_row += 1
            ws.cell(row=cp_row, column=1, value=tag)
            ws.cell(row=cp_row, column=2, value=round(cp_x,4))
            ws.cell(row=cp_row, column=3, value=round(phi2,4))
            ws.cell(row=cp_row, column=4, value=round(phi1,4))
            ws.cell(row=cp_row, column=5, value=round(phi2,4))
            ws.cell(row=cp_row, column=6, value=round(phi3,4))
            ws.cell(row=cp_row, column=7, value=f'critical_{tag}.csv')

# Data counts summary
cnt_row = cp_row + 3
ws.cell(row=cnt_row, column=1, value='Data Counts:').font = Font(bold=True)
cnt_row += 1
ws.cell(row=cnt_row, column=1, value='System').font = Font(bold=True)
ws.cell(row=cnt_row, column=2, value='Binodal pts').font = Font(bold=True)
ws.cell(row=cnt_row, column=3, value='Spinodal pts').font = Font(bold=True)
for tag in all_bino:
    cnt_row += 1
    ws.cell(row=cnt_row, column=1, value=tag)
    ws.cell(row=cnt_row, column=2, value=len(all_bino[tag]))
    ws.cell(row=cnt_row, column=3, value=len(all_spin[tag]))

ws.freeze_panes = 'A5'
wb.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
for tag in all_bino:
    print(f"  {tag}: Bino={len(all_bino[tag])}  Spin={len(all_spin[tag])}")
